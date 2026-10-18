"""V0.2 验证脚本：5 种格式解析、扫描件报错、旧格式提示、切片、上传接口入库。

运行（在 work 目录下，用 venv 解释器）：
    venv\\Scripts\\python.exe scripts\\verify_v02.py
使用临时数据库，不污染 data/kb.db。样例文件在内存中生成。
"""
import io
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
_tmp = tempfile.mkdtemp(prefix="kb_verify2_")
os.environ["DB_PATH"] = os.path.join(_tmp, "kb.db")

from app import parsers  # noqa: E402

passed, failed = 0, 0


def check(name, cond, detail=""):
    global passed, failed
    if cond:
        passed += 1
        print(f"[PASS] {name}")
    else:
        failed += 1
        print(f"[FAIL] {name} {detail}")


# ---------- 构造样例 ----------
def make_docx() -> bytes:
    import docx
    d = docx.Document()
    d.add_paragraph("向量数据库选型笔记")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "名称", "特点"
    t.cell(1, 0).text, t.cell(1, 1).text = "ChromaDB", "轻量"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def make_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "预算"
    ws.append(["项目", "金额"])
    ws.append(["服务器", 100])
    ws2 = wb.create_sheet("人员")
    ws2.append(["姓名", "张三"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def make_pdf(with_text: bool) -> bytes:
    """手写最小 PDF：with_text=False 时页面无文本层，模拟扫描件。"""
    content = b"BT /F1 18 Tf 50 700 Td (Hello Knowledge Base) Tj ET" if with_text else b""
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offs = []
    for i, o in enumerate(objs, start=1):
        offs.append(len(out))
        out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    for off in offs:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objs) + 1, xref)
    return out


# ---------- 解析器 ----------
title, text = parsers.parse_file("笔记.txt", "中文内容测试".encode("gbk"))
check("TXT(GBK) 解析且标题取文件名", title == "笔记" and "中文内容" in text, text)
title, text = parsers.parse_file("a.md", "# 标题\n- 列表项".encode("utf-8"))
check("MD 解析非空", "列表项" in text)
title, text = parsers.parse_file("选型.docx", make_docx())
check("DOCX 含段落与表格", "向量数据库选型" in text and "ChromaDB" in text, text)
title, text = parsers.parse_file("预算.xlsx", make_xlsx())
check("XLSX 多 sheet 按表名分区", "工作表：预算" in text and "工作表：人员" in text and "服务器\t100" in text, text)
title, text = parsers.parse_file("t.pdf", make_pdf(True))
check("PDF 文本层提取", "Hello Knowledge Base" in text, text)
try:
    parsers.parse_file("scan.pdf", make_pdf(False))
    check("扫描版 PDF 给出提示", False, "未抛错")
except parsers.ParseError as e:
    check("扫描版 PDF 给出提示不崩溃", "无法提取文字" in str(e), str(e))
try:
    parsers.parse_file("old.doc", b"x")
    check("旧版 .doc 给提示", False)
except parsers.ParseError as e:
    check("旧版 .doc 给提示", ".docx" in str(e), str(e))
try:
    parsers.parse_file("bad.docx", b"not a zip")
    check("损坏 DOCX 报错不崩溃", False)
except parsers.ParseError:
    check("损坏 DOCX 报错不崩溃", True)
try:
    parsers.parse_file("empty.txt", b"   \n ")
    check("空文本拒绝入库", False)
except parsers.ParseError:
    check("空文本拒绝入库", True)

# ---------- 切片 ----------
long_text = "\n".join(f"第{i}段。" + "测试内容很长。" * 30 for i in range(40))
chunks = parsers.split_chunks(long_text, size=500, overlap=50)
check("长文本被切成多片", len(chunks) > 1, str(len(chunks)))
check("每片不超过 size+overlap+段长容差", all(len(c[2]) <= 500 + 50 + 1 for c in chunks),
      str(max(len(c[2]) for c in chunks)))
check("切片覆盖全部段落(首尾)", "第0段" in chunks[0][2] and "第39段" in chunks[-1][2])
check("相邻切片有重叠", chunks[0][2][-20:] in chunks[1][2] or chunks[1][2][:20] in chunks[0][2])
check("短文本只一片", len(parsers.split_chunks("短文本")) == 1)
check("空文本无切片", parsers.split_chunks("  ") == [])

# ---------- 上传接口（需 main.py 接入 parsers）----------
try:
    from fastapi.testclient import TestClient
    from app.main import app
    client = TestClient(app)

    def up(name, data):
        return client.post("/api/v1/documents/upload", files={"file": (name, data)})

    for name, data in [("a.docx", make_docx()), ("b.xlsx", make_xlsx()), ("c.pdf", make_pdf(True))]:
        r = up(name, data)
        check(f"上传 {name} 入库", r.status_code == 200 and r.json()["data"].get("doc_id"), r.text)
    r = up("scan.pdf", make_pdf(False))
    check("上传扫描 PDF 返回 400 且提示", r.status_code == 400 and "无法提取文字" in r.text, r.text)
    r = client.post("/api/v1/search", json={"query": "ChromaDB"})
    check("上传的 DOCX 可被关键词检索", r.json()["data"]["results"] != [], r.text)
except Exception as e:  # noqa: BLE001
    check("上传接口集成", False, repr(e))

print(f"\n通过 {passed} / 失败 {failed}")
sys.exit(1 if failed else 0)
